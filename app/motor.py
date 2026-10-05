"""Motor de datos de la app de control logístico (SQLite + reglas de negocio).

Modelo: reembarque -> remision (file aduana) -> sae, y una bitácora de eventos
(hitos con fecha-hora real). Fase, semáforo y KPI se calculan desde los eventos.
"""
import io
import re
import sqlite3
import statistics
import threading
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import sys

# Junto al .exe si la app está empaquetada con PyInstaller; si no, junto a este archivo.
_BASE = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).parent
DB_PATH = _BASE / "datos" / "logistica.db"
_lock = threading.RLock()

# ------------------------------------------------------------------ catálogos
# HITO_ID: (orden, paso 1-10, fase, nivel, dueño, próxima acción, ANS de la fase en horas)
HITOS = {
    "H000": (0, 1, "F00 Remisión recibida", "REMISION", "AA KN", "Enviar borrador SAE", 4),
    "H010": (10, 2, "F10 Clasificación solicitada", "REMISION", "CLIENTE", "Responder clasificación", 3),
    "H015": (15, 2, "F10 Clasificación respondida", "REMISION", "AA KN", "Enviar borrador SAE", 4),
    "H020": (20, 3, "F20 Borrador SAE enviado", "REMISION", "KN CDLI", "Aprobar SAE en bodega", 24),
    "H030": (30, 4, "F30 Aprobado por bodega", "REMISION", "AA KN", "Autorizar SAE", 4),
    "H040": (40, 5, "F40 SAE autorizada", "SAE", "TERMINAL", "Ingreso a puerto", 48),
    "H045": (45, 6, "F45 Inspección física", "SAE", "DIAN", "Liberar inspección", 48),
    "H046": (46, 6, "F45 Inspección liberada", "SAE", "TERMINAL", "Autorizar embarque", 24),
    "H050": (50, 6, "F50 Autoriza embarque", "REEMBARQUE", "NAVIERA", "Zarpe / manifiesto", 72),
    "H060": (60, 7, "F60 Embarcado (manifiesto)", "SAE", "AA KN", "Presentar DEX", 72),
    "H070": (70, 8, "F70 DEX presentada", "SAE", "AA KN", "Cerrar carpeta documental", 120),
    "H080": (80, 9, "F80 Cierre documental", "REMISION", "ADMIN", "Prefacturar", 240),
    "H090": (90, 10, "F90 Facturado", "REMISION", "—", "—", 0),
}
PASOS = ["Recibida", "Clasif.", "Borrador", "Bodega", "SAE", "Puerto", "Embarque", "DEX", "Cierre", "Factura"]

CAUSALES = {
    "REM01": ("Remisión incompleta", "CLIENTE"), "REM02": ("Remisión con datos errados", "CLIENTE"),
    "REM03": ("Cambio de remisión", "CLIENTE"), "REM04": ("Recibida fuera de horario / fin de semana", "CLIENTE"),
    "CLA01": ("Model code sin clasificar", "AA KN"), "CLA02": ("Respuesta tardía del cliente a clasificación", "CLIENTE"),
    "CLA03": ("PA sugerida difiere", "CLIENTE"),
    "SAE01": ("Demora en elaboración SAE", "AA KN"), "SAE02": ("Demora en elaboración SAE por clasificación", "AA KN"),
    "SAE03": ("Demora en reserva / cambio de transportador", "TERCERO"),
    "SAE04": ("Remisión para adelantar sin información completa", "KN CDLI"),
    "SAE05": ("Demora +3 h en aprobación por bodega", "KN CDLI"), "SAE06": ("Demora +6 h en aprobación por bodega", "KN CDLI"),
    "SAE07": ("Demora +24 h en aprobación por bodega", "KN CDLI"),
    "SAE08": ("Error o falta de información en la remisión", "KN CDLI"), "SAE09": ("Operación fin de semana", "KN CDLI"),
    "SAE10": ("Rechazo DIAN de SAE", "AA KN"), "SAE11": ("Error de pesos", "KN CDLI"), "SAE12": ("Error de bultos", "KN CDLI"),
    "SAE13": ("Error de valor FOB", "AA KN"), "SAE14": ("Error de ítems / cantidades", "AA KN"),
    "PTO01": ("Inspección física", "DIAN"), "PTO02": ("Cierre de terminal", "TERMINAL"), "PTO03": ("Roll-over de naviera", "NAVIERA"),
    "DEX01": ("DEX manual por contingencia", "DIAN"), "DEX02": ("Corrección de DEX", "AA KN"),
    "DEX03": ("Diferencia manifiesto vs SAE", "AA KN"),
}
TRANSPORTISTAS = {  # id: (nombre, patrones de alias normalizados)
    "MAEU": ("MAERSK", ["MAERSK"]), "HLCU": ("HAPAG-LLOYD", ["HAPAG"]), "ZIMU": ("ZIM", ["ZIM"]),
    "ONEY": ("OCEAN NETWORK EXPRESS", ["OCEAN NETWORK", "ONE "]), "CMDU": ("CMA CGM", ["CMA CGM"]),
    "EGLV": ("EVERGREEN", ["EVERGREEN"]), "KN": ("KUEHNE+NAGEL", ["KUEHNE"]),
    "L360": ("LOGÍSTICA 360", ["LOGISTICA 36"]), "DHL": ("DHL", ["DHL"]), "COPA": ("COPA AIRLINES CARGO", ["COPA"]),
}
PAISES = {"CL": "Chile", "MX": "México", "UY": "Uruguay", "AR": "Argentina", "PA": "Panamá", "CR": "Costa Rica",
          "PY": "Paraguay", "SV": "El Salvador", "IE": "Irlanda", "US": "Estados Unidos", "CA": "Canadá",
          "CO": "Colombia", "PE": "Perú", "EC": "Ecuador", "GT": "Guatemala", "BR": "Brasil", "ES": "España"}
PAIS_NOMBRE = {"CHILE": "CL", "MEXICO": "MX", "URUGUAY": "UY", "ARGENTINA": "AR", "PANAMA": "PA", "COSTA RICA": "CR",
               "PARAGUAY": "PY", "EL SALVADOR": "SV", "SALVADOR": "SV", "IRLANDA": "IE", "ESTADOS UNIDOS": "US",
               "CANADA": "CA", "COLOMBIA": "CO", "PERU": "PE", "ECUADOR": "EC", "GUATEMALA": "GT", "BRASIL": "BR"}
LISTAS = {
    "terminal": ["CONTECAR", "SPRC"], "modo": ["MAR", "AER", "TER"],
    "canal": ["AUTORIZA_EMBARQUE", "INSPECCION_FISICA", "INSPECCION_DOCUMENTAL"],
    "dex_tipo": ["AUTOMATICA", "MANUAL", "CORRECCION"],
    "revision": ["EN_DIGITACION", "EN_REVISION", "PDTE_VALORES", "CONFORME"],
    "gestor": ["CARPETA_CREADA", "PDTE", "FALTA_DEX", "COMPLETO"],
    "fuente": ["MANUAL", "CORREO", "MUISCA", "WMS", "NAVIERA", "IMPORTACION"],
    "proyecto": ["BOOSTER", "RT", "DOLLARCITY"],
}
PARAMS_DEF = {
    "ANS_ELAB_SAE_H": (4, "Horas recepción → borrador SAE"),
    "ANS_TOTAL_SAE_H": (28, "Horas recepción → SAE autorizada"),
    "ANS_CLASIF_H": (3, "Horas solicitud → respuesta de clasificación"),
    "TOL_PESO_KG": (0.5, "Tolerancia kg peso SAE vs manifestado"),
    "UMBRAL_AMBAR": (0.75, "Fracción del ANS de fase para pasar a ámbar"),
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS reembarque (
  id TEXT PRIMARY KEY, proyecto TEXT, modo TEXT, agente TEXT, naviera TEXT, booking TEXT,
  terminal TEXT, pol TEXT DEFAULT 'COCTG', pod TEXT, contenedor TEXT, bl TEXT, buque TEXT, etd TEXT);
CREATE TABLE IF NOT EXISTS remision (
  file TEXT PRIMARY KEY, remision TEXT, reembarque TEXT, pais TEXT, destinatario TEXT,
  packing_list TEXT, digitador TEXT, apoyo TEXT, revision TEXT, gestor TEXT,
  bloqueo INTEGER DEFAULT 0, bloqueo_causal TEXT, observacion TEXT);
CREATE TABLE IF NOT EXISTS sae (
  id TEXT PRIMARY KEY, file TEXT, cantidad REAL, fob_usd REAL, peso_neto REAL, peso_bruto REAL,
  bultos INTEGER, peso_manif REAL, bultos_manif INTEGER, model_codes INTEGER, partidas INTEGER,
  canal TEXT, dex TEXT, dex_tipo TEXT, sae_reemplazada TEXT, causal TEXT);
CREATE TABLE IF NOT EXISTS evento (
  id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT NOT NULL, nivel TEXT, reembarque TEXT,
  file TEXT, sae TEXT, hito TEXT NOT NULL, responsable TEXT, fuente TEXT, causal TEXT,
  nota TEXT, estimado INTEGER DEFAULT 0);
CREATE INDEX IF NOT EXISTS ix_ev_file ON evento(file);
CREATE INDEX IF NOT EXISTS ix_ev_reemb ON evento(reembarque);
CREATE INDEX IF NOT EXISTS ix_sae_file ON sae(file);
CREATE TABLE IF NOT EXISTS persona (id TEXT PRIMARY KEY, nombre TEXT UNIQUE, rol TEXT);
CREATE TABLE IF NOT EXISTS parametro (clave TEXT PRIMARY KEY, valor REAL);
"""

FMT = "%Y-%m-%d %H:%M"


def conectar():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH, check_same_thread=False)
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA)
    return con


def filas(con, sql, args=()):
    return [dict(r) for r in con.execute(sql, args)]


def params(con):
    p = {k: v for k, (v, _) in PARAMS_DEF.items()}
    for r in con.execute("SELECT clave, valor FROM parametro"):
        p[r["clave"]] = r["valor"]
    return p


def meta(con):
    return {
        "hitos": [{"id": k, "orden": v[0], "paso": v[1], "fase": v[2], "nivel": v[3], "dueno": v[4],
                   "accion": v[5], "ans": v[6]} for k, v in HITOS.items()],
        "pasos": PASOS,
        "causales": [{"id": k, "desc": d, "resp": r} for k, (d, r) in CAUSALES.items()],
        "transportistas": [{"id": k, "nombre": v[0]} for k, v in TRANSPORTISTAS.items()],
        "paises": [{"id": k, "nombre": v} for k, v in PAISES.items()],
        "personas": filas(con, "SELECT * FROM persona ORDER BY nombre"),
        "reembarques": [r["id"] for r in con.execute("SELECT id FROM reembarque ORDER BY id DESC")],
        "listas": LISTAS,
        "params": [{"clave": k, "valor": params(con)[k], "desc": d} for k, (_, d) in PARAMS_DEF.items()],
        "vacia": con.execute("SELECT COUNT(*) FROM remision").fetchone()[0] == 0,
    }


# ------------------------------------------------------------------ cálculo
def _dt(s):
    return datetime.strptime(s, FMT) if s else None


def _horas(a, b):
    return round((b - a).total_seconds() / 3600, 1)


def calcular(con, ahora=None):
    """Devuelve (remisiones_calculadas, saes_calculadas) con fase, semáforo y KPI."""
    ahora = ahora or datetime.now()
    p = params(con)
    rems = filas(con, "SELECT r.*, b.terminal, b.naviera, b.agente, b.contenedor, b.pod FROM remision r "
                      "LEFT JOIN reembarque b ON b.id = r.reembarque")
    saes = filas(con, "SELECT * FROM sae")
    evs = filas(con, "SELECT * FROM evento")
    personas = {r["id"]: r["nombre"] for r in con.execute("SELECT id, nombre FROM persona")}

    ev_file, ev_reemb, ev_sae = defaultdict(list), defaultdict(list), defaultdict(list)
    for e in evs:
        e["_dt"] = _dt(e["fecha"])
        if e["file"]:
            ev_file[e["file"]].append(e)
        elif e["reembarque"]:
            ev_reemb[e["reembarque"]].append(e)
        if e["sae"]:
            ev_sae[e["sae"]].append(e)

    def primero(lst, hito, real=False):
        ds = [e["_dt"] for e in lst if e["hito"] == hito and not (real and e["estimado"])]
        return min(ds) if ds else None

    sae_por_file = defaultdict(list)
    for s in saes:
        f_ev = ev_file.get(s["file"], [])
        recep = primero(f_ev, "H000")
        borr = primero(f_ev, "H020")
        sae_ok = primero(ev_sae.get(s["id"], []), "H040", real=True)
        sae_est = any(e["estimado"] for e in ev_sae.get(s["id"], []) if e["hito"] == "H040")
        s["h_elab"] = _horas(recep, borr) if recep and borr else None
        s["h_total"] = _horas(recep, sae_ok) if recep and sae_ok else None
        s["ans_elab"] = "PENDIENTE" if s["h_elab"] is None else ("CUMPLE" if s["h_elab"] <= p["ANS_ELAB_SAE_H"] else "NO_CUMPLE")
        if sae_est:
            s["ans_sae"] = "ESTIMADO"
        else:
            s["ans_sae"] = "PENDIENTE" if s["h_total"] is None else ("CUMPLE" if s["h_total"] <= p["ANS_TOTAL_SAE_H"] else "NO_CUMPLE")
        if s["peso_manif"] is None and s["bultos_manif"] is None:
            s["dif_peso"], s["conciliado"] = None, None
        else:
            s["dif_peso"] = round((s["peso_manif"] or 0) - (s["peso_bruto"] or 0), 2)
            s["conciliado"] = abs(s["dif_peso"]) <= p["TOL_PESO_KG"] and (s["bultos_manif"] == s["bultos"])
        s["falta_causal"] = (s["ans_sae"] == "NO_CUMPLE" or s["ans_elab"] == "NO_CUMPLE" or s["conciliado"] is False) and not s["causal"]
        s["dex_dt"] = primero(ev_sae.get(s["id"], []), "H070")
        sae_por_file[s["file"]].append(s)

    for r in rems:
        lst = ev_file.get(r["file"], []) + ev_reemb.get(r["reembarque"], [])
        r["saes"] = len(sae_por_file[r["file"]])
        r["peso_bruto"] = round(sum(s["peso_bruto"] or 0 for s in sae_por_file[r["file"]]), 2)
        r["fob_usd"] = round(sum(s["fob_usd"] or 0 for s in sae_por_file[r["file"]]), 2)
        r["no_conciliadas"] = sum(1 for s in sae_por_file[r["file"]] if s["conciliado"] is False)
        r["sin_causal"] = sum(1 for s in sae_por_file[r["file"]] if s["falta_causal"])
        r["inspeccion"] = any(s["canal"] == "INSPECCION_FISICA" for s in sae_por_file[r["file"]])
        rec = primero(lst, "H000")
        r["recepcion"] = rec.strftime(FMT) if rec else None
        r["digitador_nombre"] = personas.get(r["digitador"], r["digitador"])
        if not lst:
            r.update(orden=None, fase="Sin hitos", paso=0, dueno="", accion="Registrar recepción",
                     ultimo=None, horas=None, ans=None, semaforo="AMBAR")
            continue
        orden = max(HITOS[e["hito"]][0] for e in lst if e["hito"] in HITOS)
        hid = next(k for k, v in HITOS.items() if v[0] == orden)
        o, paso, fase, _, dueno, accion, ans = HITOS[hid]
        ult = max(e["_dt"] for e in lst if e["hito"] == hid)
        horas = _horas(ult, ahora)
        if orden >= 80:
            sem = "CERRADO"
        elif r["bloqueo"] or r["no_conciliadas"] or orden == 45:
            sem = "ROJO"
        elif ans and horas > ans:
            sem = "ROJO"
        elif ans and horas > ans * p["UMBRAL_AMBAR"]:
            sem = "AMBAR"
        else:
            sem = "VERDE"
        r.update(orden=orden, fase=fase, paso=paso, dueno=dueno, accion=accion,
                 ultimo=ult.strftime(FMT), horas=horas, ans=ans, semaforo=sem)
    return rems, saes


SEV = {"ROJO": 0, "AMBAR": 1, "VERDE": 2, "CERRADO": 3}


def listar_remisiones(con, q="", fase="", semaforo="", pais="", digitador="", estado="abiertos", limite=400):
    rems, _ = calcular(con)
    qn = _norm(q)
    out = []
    for r in rems:
        if estado == "abiertos" and r["semaforo"] == "CERRADO":
            continue
        if fase and not (r["fase"] or "").startswith(fase):
            continue
        if semaforo and r["semaforo"] != semaforo:
            continue
        if pais and r["pais"] != pais:
            continue
        if digitador and r["digitador"] != digitador:
            continue
        if qn and qn not in _norm(" ".join(str(r.get(k) or "") for k in
                                          ("file", "remision", "reembarque", "destinatario", "packing_list", "contenedor"))):
            continue
        out.append(r)
    out.sort(key=lambda r: (SEV.get(r["semaforo"], 9), -(r["horas"] or 0)))
    return {"total": len(out), "filas": out[:limite]}


def detalle(con, file):
    rems, saes = calcular(con)
    r = next((x for x in rems if x["file"] == file), None)
    if not r:
        return None
    r["saes_det"] = [s | {"dex_dt": s["dex_dt"].strftime(FMT) if s["dex_dt"] else None}
                     for s in saes if s["file"] == file]
    r["eventos"] = filas(con, "SELECT * FROM evento WHERE file = ? OR (file IS NULL AND reembarque = ?) "
                              "ORDER BY fecha, id", (file, r["reembarque"]))
    for e in r["eventos"]:
        h = HITOS.get(e["hito"])
        e["fase"] = h[2] if h else e["hito"]
    r["reembarque_det"] = dict(con.execute("SELECT * FROM reembarque WHERE id = ?", (r["reembarque"],)).fetchone() or {})
    return r


def tablero(con, desde=None, hasta=None):
    rems, saes = calcular(con)
    p = params(con)

    def en_rango(fecha):
        if not fecha:
            return not (desde or hasta)
        return (not desde or fecha[:10] >= desde) and (not hasta or fecha[:10] <= hasta)

    rems_r = [r for r in rems if en_rango(r["recepcion"])]
    files_r = {r["file"] for r in rems_r}
    saes_r = [s for s in saes if s["file"] in files_r]
    abiertos = [r for r in rems if r["semaforo"] != "CERRADO"]

    def pct(lst, campo):
        c = sum(1 for s in lst if s[campo] == "CUMPLE")
        n = sum(1 for s in lst if s[campo] in ("CUMPLE", "NO_CUMPLE"))
        return round(c / n, 4) if n else None

    h_tot = [s["h_total"] for s in saes_r if s["h_total"] is not None and s["ans_sae"] != "ESTIMADO"]
    por_fase = defaultdict(Counter)
    for r in abiertos:
        por_fase[r["fase"]][r["semaforo"]] += 1
    orden_fase = {v[2]: v[0] for v in HITOS.values()}
    fases = [{"fase": f, **por_fase[f]} for f in sorted(por_fase, key=lambda f: orden_fase.get(f, -1))]

    mes = defaultdict(lambda: {"files": 0, "saes": 0, "cumple": 0, "medibles": 0, "estimadas": 0})
    file_mes = {}
    for r in rems_r:
        if r["recepcion"]:
            m = r["recepcion"][:7]
            mes[m]["files"] += 1
            file_mes[r["file"]] = m
    for s in saes_r:
        m = file_mes.get(s["file"])
        if not m:
            continue
        mes[m]["saes"] += 1
        if s["ans_sae"] in ("CUMPLE", "NO_CUMPLE"):
            mes[m]["medibles"] += 1
            mes[m]["cumple"] += s["ans_sae"] == "CUMPLE"
        elif s["ans_sae"] == "ESTIMADO":
            mes[m]["estimadas"] += 1
    meses = [{"mes": k, **v, "pct": round(v["cumple"] / v["medibles"], 4) if v["medibles"] else None}
             for k, v in sorted(mes.items())][-18:]

    caus = Counter()
    for s in saes_r:
        if s["causal"]:
            caus[s["causal"]] += 1
    evc = con.execute("SELECT causal, file, fecha FROM evento WHERE causal IS NOT NULL AND causal <> ''").fetchall()
    for e in evc:
        if e["file"] in files_r:
            caus[e["causal"]] += 1
    causales = [{"id": k, "desc": CAUSALES.get(k, (k, ""))[0], "resp": CAUSALES.get(k, ("", "?"))[1], "n": n}
                for k, n in caus.most_common(10)]

    personas = {r["id"]: r["nombre"] for r in con.execute("SELECT id, nombre FROM persona")}
    dig = Counter()
    file_dig = {r["file"]: r["digitador"] for r in rems_r}
    for s in saes_r:
        dig[file_dig.get(s["file"])] += 1
    digitadores = [{"id": k, "nombre": personas.get(k, k or "Sin asignar"), "n": n} for k, n in dig.most_common(8)]
    paises = Counter(r["pais"] for r in rems_r)

    return {
        "tiles": {
            "abiertos": len(abiertos),
            "rojo": sum(1 for r in abiertos if r["semaforo"] == "ROJO"),
            "ambar": sum(1 for r in abiertos if r["semaforo"] == "AMBAR"),
            "files": len(rems_r), "saes": len(saes_r),
            "pct_elab": pct(saes_r, "ans_elab"), "pct_sae": pct(saes_r, "ans_sae"),
            "mediana_h_sae": round(statistics.median(h_tot), 1) if h_tot else None,
            "no_conciliadas": sum(1 for s in saes_r if s["conciliado"] is False),
            "sin_causal": sum(1 for s in saes_r if s["falta_causal"]),
            "estimadas": sum(1 for s in saes_r if s["ans_sae"] == "ESTIMADO"),
            "fob_usd": round(sum(s["fob_usd"] or 0 for s in saes_r), 2),
            "ans_sae_h": p["ANS_TOTAL_SAE_H"], "ans_elab_h": p["ANS_ELAB_SAE_H"],
        },
        "fases": fases, "meses": meses, "causales": causales, "digitadores": digitadores,
        "paises": [{"id": k or "?", "nombre": PAISES.get(k, k or "Sin país"), "n": n} for k, n in paises.most_common(10)],
    }


# ------------------------------------------------------------------ escritura
def _limpia(d, campos):
    out = {}
    for c in campos:
        v = d.get(c)
        if isinstance(v, str):
            v = v.strip() or None
        out[c] = v
    return out


def guardar(con, tabla, d):
    campos = {"reembarque": ["id", "proyecto", "modo", "agente", "naviera", "booking", "terminal", "pol", "pod",
                             "contenedor", "bl", "buque", "etd"],
              "remision": ["file", "remision", "reembarque", "pais", "destinatario", "packing_list", "digitador",
                           "apoyo", "revision", "gestor", "bloqueo", "bloqueo_causal", "observacion"],
              "sae": ["id", "file", "cantidad", "fob_usd", "peso_neto", "peso_bruto", "bultos", "peso_manif",
                      "bultos_manif", "model_codes", "partidas", "canal", "dex", "dex_tipo", "sae_reemplazada", "causal"]}[tabla]
    v = _limpia(d, campos)
    clave = campos[0]
    if not v[clave]:
        raise ValueError(f"Falta {clave}")
    if tabla == "sae":
        if not re.fullmatch(r"\d{13}", v["id"]):
            raise ValueError("La SAE debe tener 13 dígitos")
        if v["dex"] and not re.fullmatch(r"\d{13}", v["dex"]):
            raise ValueError("La DEX debe tener 13 dígitos")
        if v["peso_neto"] and v["peso_bruto"] and float(v["peso_neto"]) > float(v["peso_bruto"]):
            raise ValueError("El peso neto no puede superar el bruto")
    if tabla == "reembarque" and v["contenedor"] and not re.fullmatch(r"[A-Z]{4}\d{7}", v["contenedor"].upper()):
        raise ValueError("Contenedor ISO 6346: 4 letras + 7 dígitos")
    if tabla == "remision":
        v["bloqueo"] = 1 if v["bloqueo"] in (1, True, "1", "S", "true") else 0
    with _lock:
        con.execute(f"INSERT OR REPLACE INTO {tabla} ({','.join(campos)}) VALUES ({','.join('?' * len(campos))})",
                    [v[c] for c in campos])
        con.commit()
    return v


def guardar_evento(con, d):
    v = _limpia(d, ["fecha", "nivel", "reembarque", "file", "sae", "hito", "responsable", "fuente", "causal", "nota"])
    if v["hito"] not in HITOS:
        raise ValueError("Hito inválido")
    fecha = (v["fecha"] or "").replace("T", " ")[:16]
    try:
        datetime.strptime(fecha, FMT)
    except ValueError:
        raise ValueError("Fecha-hora inválida (yyyy-mm-dd hh:mm)")
    if datetime.strptime(fecha, FMT) > datetime.now() + timedelta(hours=1):
        raise ValueError("No se registran hitos en el futuro: la fecha debe ser la real del hecho")
    v["fecha"] = fecha
    if v["file"] and not v["reembarque"]:
        row = con.execute("SELECT reembarque FROM remision WHERE file = ?", (v["file"],)).fetchone()
        v["reembarque"] = row["reembarque"] if row else None
    if v["sae"] and not v["file"]:
        row = con.execute("SELECT file FROM sae WHERE id = ?", (v["sae"],)).fetchone()
        v["file"] = row["file"] if row else None
    v["nivel"] = "SAE" if v["sae"] else ("REMISION" if v["file"] else "REEMBARQUE")
    if not (v["file"] or v["reembarque"]):
        raise ValueError("Indique el file o el reembarque")
    with _lock:
        cur = con.execute("INSERT INTO evento (fecha, nivel, reembarque, file, sae, hito, responsable, fuente, causal, nota) "
                          "VALUES (?,?,?,?,?,?,?,?,?,?)",
                          [v[k] for k in ("fecha", "nivel", "reembarque", "file", "sae", "hito", "responsable",
                                          "fuente", "causal", "nota")])
        con.commit()
    return {"id": cur.lastrowid}


def borrar_evento(con, eid):
    with _lock:
        con.execute("DELETE FROM evento WHERE id = ?", (eid,))
        con.commit()


def guardar_params(con, d):
    with _lock:
        for k in PARAMS_DEF:
            if k in d and d[k] not in (None, ""):
                con.execute("INSERT OR REPLACE INTO parametro VALUES (?, ?)", (k, float(d[k])))
        con.commit()


def guardar_persona(con, nombre, rol="DIGITADOR"):
    nombre = _norm(nombre)
    if not nombre:
        return None
    row = con.execute("SELECT id FROM persona WHERE nombre = ?", (nombre,)).fetchone()
    if row:
        return row["id"]
    n = con.execute("SELECT COUNT(*) FROM persona").fetchone()[0] + 1
    pid = f"P{n:03d}"
    while con.execute("SELECT 1 FROM persona WHERE id = ?", (pid,)).fetchone():
        n += 1
        pid = f"P{n:03d}"
    con.execute("INSERT INTO persona VALUES (?, ?, ?)", (pid, nombre, rol))
    return pid


def vaciar(con):
    with _lock:
        for t in ("evento", "sae", "remision", "reembarque", "persona"):
            con.execute(f"DELETE FROM {t}")
        con.commit()


# ------------------------------------------------------------------ demo
def cargar_demo(con):
    """Datos ficticios relativos a 'ahora' para ver el tablero vivo."""
    ahora = datetime.now().replace(second=0, microsecond=0)
    t = lambda h: (ahora - timedelta(hours=h)).strftime(FMT)
    with _lock:
        for nombre in ("ANALISTA UNO", "ANALISTA DOS", "ANALISTA TRES", "RESPONSABLE DEX"):
            guardar_persona(con, nombre)
        con.commit()
    P1, P2, P3, P4 = "P001", "P002", "P003", "P004"
    guardar(con, "reembarque", dict(id="BST3-0412", proyecto="BOOSTER", modo="MAR", agente="MAEU", naviera="MAEU",
                                    booking="268114502", terminal="CONTECAR", pod="CLVAP", contenedor="MSKU7340215",
                                    bl="MAEU268114502", buque="DEMO / 640S", etd=t(10)[:10]))
    guardar(con, "reembarque", dict(id="BST3-0413", proyecto="BOOSTER", modo="MAR", agente="L360", naviera="HLCU",
                                    booking="HLCU-CTG-88231", terminal="SPRC", pod="MXZLO", contenedor="HLXU8812043"))
    guardar(con, "reembarque", dict(id="BST3-0414", proyecto="BOOSTER", modo="MAR", agente="MAEU", naviera="ZIMU",
                                    booking="ZIM-55120", terminal="CONTECAR", pod="UYMVD", contenedor="ZCSU5512046"))
    rems = [
        ("DEMO-0223-412", "BOOST-412.1-CL", "BST3-0412", "CL", "DECATHLON CHILE SPA", "2900230011", P2, "CONFORME", "PDTE"),
        ("DEMO-0223-413", "BOOST-412.2-CL", "BST3-0412", "CL", "DECATHLON CHILE SPA", "2900230012", P1, "CONFORME", "FALTA_DEX"),
        ("DEMO-0223-414", "BOOST-413.1-MX", "BST3-0413", "MX", "ARTICULOS DEPORTIVOS DECATHLON SA DE CV", "2900230020", P3, "EN_REVISION", "PDTE"),
        ("DEMO-0223-415", "BOOST-414.1-UY", "BST3-0414", "UY", "DEPORSUR SA", "2900230031", P1, "EN_DIGITACION", "CARPETA_CREADA"),
        ("DEMO-0223-416", "BOOST-414.2-UY", "BST3-0414", "UY", "DEPORSUR SA", "2900230032", P2, "EN_DIGITACION", "CARPETA_CREADA"),
        ("DEMO-0223-401", "BOOST-401.1-CL", "BST3-0412", "CL", "DECATHLON CHILE SPA", "2900229001", P3, "CONFORME", "COMPLETO"),
    ]
    for f, rm, rb, pa, de, pl, dg, rv, gs in rems:
        guardar(con, "remision", dict(file=f, remision=rm, reembarque=rb, pais=pa, destinatario=de, packing_list=pl,
                                      digitador=dg, revision=rv, gestor=gs, bloqueo=0))
    saes = [
        ("6027770000011", "DEMO-0223-412", 3647, 76754.29, 2102.55, 2486.79, 212, 2486.79, 212, 310, 74, "AUTORIZA_EMBARQUE", "6007770000101", "AUTOMATICA", None),
        ("6027770000028", "DEMO-0223-412", 6692, 121983.36, 4934.10, 5724.23, 418, 5724.23, 418, 453, 81, "AUTORIZA_EMBARQUE", "6007770000118", "AUTOMATICA", None),
        ("6027770000035", "DEMO-0223-413", 1347, 33083.82, 1462.20, 1708.47, 192, 1702.10, 190, 74, 69, "INSPECCION_FISICA", None, None, "SAE08"),
        ("6027770000042", "DEMO-0223-414", 288, 13878.52, 944.30, 1095.41, 136, None, None, 68, 73, None, None, None, None),
        ("6027770000059", "DEMO-0223-415", 4168, 54960.50, 3820.00, 4263.49, 331, None, None, 346, 77, None, None, None, None),
        ("6027770000066", "DEMO-0223-401", 4501, 112344.65, 3990.10, 4432.59, 395, 4432.59, 395, 490, 80, "AUTORIZA_EMBARQUE", "6007770000040", "AUTOMATICA", None),
    ]
    for s in saes:
        guardar(con, "sae", dict(zip(["id", "file", "cantidad", "fob_usd", "peso_neto", "peso_bruto", "bultos", "peso_manif",
                                      "bultos_manif", "model_codes", "partidas", "canal", "dex", "dex_tipo", "causal"], s)))
    ev = [  # (horas atrás, reembarque, file, sae, hito, responsable, fuente, causal, nota)
        (120, "BST3-0412", "DEMO-0223-412", None, "H000", "CDLI", "CORREO", None, "PL 2900230011"),
        (117.4, "BST3-0412", "DEMO-0223-412", None, "H020", P2, "CORREO", None, "2 SAE"),
        (112.7, "BST3-0412", "DEMO-0223-412", None, "H030", "CDLI", "CORREO", None, None),
        (98.6, "BST3-0412", "DEMO-0223-412", "6027770000011", "H040", P2, "MUISCA", None, None),
        (98.6, "BST3-0412", "DEMO-0223-412", "6027770000028", "H040", P2, "MUISCA", None, None),
        (119, "BST3-0412", "DEMO-0223-413", None, "H000", "CDLI", "CORREO", None, None),
        (113.9, "BST3-0412", "DEMO-0223-413", None, "H020", P1, "CORREO", "SAE08", "Remisión con bultos errados"),
        (108, "BST3-0412", "DEMO-0223-413", None, "H030", "CDLI", "CORREO", None, None),
        (88.8, "BST3-0412", "DEMO-0223-413", "6027770000035", "H040", P1, "MUISCA", None, None),
        (70, "BST3-0412", "DEMO-0223-413", "6027770000035", "H045", "TERMINAL", "MUISCA", "PTO01", "Selección inspección física"),
        (45, "BST3-0412", "DEMO-0223-413", "6027770000035", "H046", "TERMINAL", "MUISCA", None, None),
        (71, "BST3-0412", None, None, "H050", "TERMINAL", "MUISCA", None, "Canal: autoriza embarque"),
        (9.3, "BST3-0412", None, None, "H060", "NAVIERA", "NAVIERA", None, "Zarpe buque/viaje 640S"),
        (9.3, "BST3-0412", "DEMO-0223-412", "6027770000011", "H070", P4, "MUISCA", None, None),
        (9.3, "BST3-0412", "DEMO-0223-412", "6027770000028", "H070", P4, "MUISCA", None, None),
        (17.6, "BST3-0413", "DEMO-0223-414", None, "H000", "CDLI", "CORREO", "REM04", "Recibida fin de semana"),
        (14.2, "BST3-0413", "DEMO-0223-414", None, "H020", P3, "CORREO", None, None),
        (3.2, "BST3-0413", "DEMO-0223-414", None, "H030", "CDLI", "CORREO", None, None),
        (2.5, "BST3-0414", "DEMO-0223-415", None, "H000", "CDLI", "CORREO", None, None),
        (1.1, "BST3-0414", "DEMO-0223-416", None, "H000", "CDLI", "CORREO", None, None),
        (0.6, "BST3-0414", "DEMO-0223-416", None, "H010", P2, "CORREO", None, "3 model codes nuevos"),
        (400, "BST3-0412", "DEMO-0223-401", None, "H000", "CDLI", "CORREO", None, None),
        (397, "BST3-0412", "DEMO-0223-401", None, "H020", P3, "CORREO", None, None),
        (390, "BST3-0412", "DEMO-0223-401", None, "H030", "CDLI", "CORREO", None, None),
        (379, "BST3-0412", "DEMO-0223-401", "6027770000066", "H040", P3, "MUISCA", None, None),
        (300, "BST3-0412", "DEMO-0223-401", "6027770000066", "H060", "NAVIERA", "NAVIERA", None, None),
        (290, "BST3-0412", "DEMO-0223-401", "6027770000066", "H070", P4, "MUISCA", None, None),
        (200, "BST3-0412", "DEMO-0223-401", None, "H080", P4, "MANUAL", None, "Carpeta completa"),
    ]
    with _lock:
        for h, rb, f, s, hito, resp, fu, ca, nota in ev:
            con.execute("INSERT INTO evento (fecha, nivel, reembarque, file, sae, hito, responsable, fuente, causal, nota) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?)",
                        (t(h), "SAE" if s else ("REMISION" if f else "REEMBARQUE"), rb, f, s, hito, resp, fu, ca, nota))
        con.commit()


# ------------------------------------------------------------------ importación del consolidado
def _norm(s):
    if s is None:
        return ""
    s = str(s).replace("\xa0", " ")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


def _fecha(v):
    """Convierte datetime, serial de Excel o texto (dd-mm-yyyy, dd.mm.yyyy, ISO) a datetime. Texto no fecha -> None."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v if v.year >= 2000 else None
    if isinstance(v, (int, float)) and 30000 < v < 60000:
        return datetime(1899, 12, 30) + timedelta(days=float(v))
    s = str(v).replace("\xa0", " ").strip()
    if re.fullmatch(r"\d{5}(\.\d+)?", s):
        return _fecha(float(s))
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})[T ](\d{2}):(\d{2})", s)
    if m:
        return datetime(*map(int, m.groups()))
    m = re.fullmatch(r"(\d{1,2})[-./,](\d{1,2})[-./,](\d{4})", s)
    if m:
        d, mo, y = map(int, m.groups())
        try:
            return datetime(y, mo, d)
        except ValueError:
            return None
    return None


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("\xa0", "").strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _codigo13(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\D", "", str(v))
    return s if len(s) == 13 else None


def _transportista(v, log):
    n = _norm(v)
    if not n or n in ("PDTE", "PENDIENTE"):
        return None
    for k, (_, pats) in TRANSPORTISTAS.items():
        if any(n.startswith(p) or p in n for p in pats):
            return k
    log["transportista_desconocido"][n] += 1
    return n


def _terminal(v):
    n = _norm(v)
    term = "CONTECAR" if n.startswith("CONTECAR") else ("SPRC" if n.startswith("SPRC") else (n or None))
    return term, ("AER" if "AEREO" in n else None)


def _pais(v):
    n = _norm(v)
    if not n:
        return None
    partes = [x.strip() for x in n.split("/")]
    for p in partes:
        if p in PAIS_NOMBRE:
            return PAIS_NOMBRE[p]
    for nombre, iso in PAIS_NOMBRE.items():
        if nombre in n:
            return iso
    return None


def _destinatario(v):
    n = _norm(v).replace(",", ".")
    n = re.sub(r"\.(?=\S)", ". ", n)
    n = re.sub(r"\s+", " ", n).strip(" .")
    n = re.sub(r"\bS\. ?A\.?(?= |$)", "SA", n)
    return re.sub(r" MEXICO$", "", n)


REVISION_MAP = {"OK CONFORMIDAD": "CONFORME", "EN REVISION": "EN_REVISION", "EN DIGITACION": "EN_DIGITACION",
                "PDTE CONFIRMAR VALORES": "PDTE_VALORES"}
GESTOR_MAP = {"DOC COMPLETOS": "COMPLETO", "FALTA DEX": "FALTA_DEX", "PDTE": "PDTE", "CARPETA CREADA": "CARPETA_CREADA"}
COLS = {  # clave interna -> encabezado normalizado del consolidado
    "modo": "MODO DE TRANSPORTE", "reembarque": "# REEMBARQUE", "remision": "REMISION NUMERO",
    "recepcion": "FECHA RECEPCION REMISION / DOCUMENTOS E INFORMACION COMPLETA", "destino": "DESTINO",
    "pais": "PAIS / CIUDAD", "agente": "AGENTE DE CARGA", "naviera": "LINEA NAVIERA", "estado": "ESTADO",
    "gestor": "ESTADO GESTOR DOCUMENTAL", "pl": "PACKING LIST", "asunto": "ASUNTO", "file": "FILE ADUANA",
    "digitador": "DIGITADOR", "lugar": "LUGAR DE INGRESO DE MERCANCIA", "sae": "SAE",
    "borrador": "FECHA ENVIO BORRADOR", "aprobacion": "FECHA APROBACION SAE (KN CL)", "fecha_sae": "FECHA SAE",
    "cumpl_sae": "CUMPLIMIENTO SAE", "causal": "CAUSAL DE INCUMPLIMIENTO", "cantidades": "CANTIDADES",
    "valor": "VALORTOTAL", "neto": "PESO NETO KG (TOTAL)", "bruto": "PESO BRUTO KG (TOTAL)", "bultos": "BULTOS",
    "model": "CANT. MODEL CODE X SAE (REF. DECATHLON)", "partidas": "CANT. PARTIDAS ARANCELARIAS X SAE",
    "peso_manif": "PESO BRUTO MANIFESTADO", "bultos_manif": "BULTOS MANIFESTADOS", "manifiesto": "FECHA MANIFIESTO",
    "dex": "DEX", "fecha_dex": "FECHA DEX", "responsable": "RESPONSABLE", "sae_reemp": "SAE INICIAL CON REEMPLAZO",
    "correccion": "CORECCION DEX", "facturacion": "FECHA FACTURACION",
    "clasif_sol": "FECHA SOLICITUD CLASIFICACION", "clasif_resp": "FECHA RESPUESTA CLIENTE CLASIFICACION",
    "obs": "OBSERVACIONES BOG",
}


def _hdr(h):
    return _norm(h).replace(",", ".").replace(" .", ".").replace(". ", ".").replace(".", ". ").replace("  ", " ").strip()


def importar_consolidado(con, contenido, hoja=None):
    """Importa la hoja DATA del CONSOLIDADO_GENERAL al modelo. Devuelve un informe de calidad."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = hoja or next((n for n in ("DATA", "Reporte") if n in wb.sheetnames), wb.sheetnames[0])
    ws = wb[hoja]
    log = {"hoja": hoja, "filas_leidas": 0, "saes": 0, "files": set(), "reembarques": set(), "eventos": 0,
           "sin_sae": 0, "fechas_estimadas": 0, "texto_en_fecha": Counter(), "fechas_texto_convertidas": 0,
           "transportista_desconocido": Counter(), "alias_unificados": Counter(), "sin_pais": 0,
           "duplicadas": 0, "no_conciliadas": 0}
    idx = None
    vacias = 0
    vistos_sae = set()
    eventos = []
    ev_file_vistos = set()
    remisiones, reembarques = {}, {}

    def ev(fecha, hito, rb, f, s=None, resp=None, causal=None, nota=None, estimado=0):
        if fecha is None:
            return
        clave = (hito, f, s, rb)
        if s is None and clave in ev_file_vistos:
            return
        ev_file_vistos.add(clave)
        eventos.append((fecha.strftime(FMT), "SAE" if s else ("REMISION" if f else "REEMBARQUE"), rb, f, s, hito,
                        resp, "IMPORTACION", causal, nota, estimado))

    with _lock:
        for fila in ws.iter_rows(values_only=True):
            if idx is None:
                hdrs = [_hdr(h) for h in fila]
                idx = {}
                for k, h in COLS.items():
                    hn = _hdr(h)
                    for i, x in enumerate(hdrs):
                        if x == hn and k not in idx:
                            idx[k] = i
                if "sae" not in idx or "file" not in idx:
                    raise ValueError(f"La hoja '{hoja}' no tiene las columnas SAE y FILE ADUANA del consolidado")
                continue
            g = lambda k: fila[idx[k]] if k in idx and idx[k] < len(fila) else None
            if not any(v not in (None, "") for v in fila[:40]):
                vacias += 1
                if vacias > 300:
                    break
                continue
            vacias = 0
            log["filas_leidas"] += 1
            sae_id = _codigo13(g("sae"))
            file = _norm(g("file")) or None
            if not sae_id or not file:
                log["sin_sae"] += 1
                continue
            if sae_id in vistos_sae:
                log["duplicadas"] += 1
            vistos_sae.add(sae_id)
            rb = _norm(g("reembarque")) or None
            # ---- reembarque
            if rb and rb not in reembarques:
                term, modo_aer = _terminal(g("lugar"))
                modo_txt = _norm(g("modo"))
                pod = re.search(r"COCTG\s*-\s*([A-Z]{5})", _norm(g("asunto")))
                reembarques[rb] = dict(id=rb, proyecto="BOOSTER" if rb.startswith("BOOST") else rb.split("-")[0],
                                       modo="AER" if (modo_aer or modo_txt.startswith("AER")) else "MAR",
                                       agente=_transportista(g("agente"), log), naviera=_transportista(g("naviera"), log),
                                       terminal=term, pod=pod.group(1) if pod else None, pol="COCTG")
            for k in ("agente", "naviera", "lugar", "destino", "digitador", "pais"):
                raw = g(k)
                if raw is not None and str(raw) != _norm(raw):
                    log["alias_unificados"][k] += 1
            # ---- remisión
            if file not in remisiones:
                dig = [x for x in re.split(r"\s*-\s*", _norm(g("digitador"))) if x and x != "PDTE"]
                pais = _pais(g("pais"))
                if not pais:
                    log["sin_pais"] += 1
                remisiones[file] = dict(
                    file=file, remision=re.sub(r"\s+", "", _norm(g("remision"))) or None, reembarque=rb, pais=pais,
                    destinatario=_destinatario(g("destino")) or None,
                    packing_list=str(g("pl")).replace("\xa0", " ").strip() if g("pl") is not None else None,
                    digitador=guardar_persona(con, dig[0]) if dig else None,
                    apoyo=guardar_persona(con, dig[1]) if len(dig) > 1 else None,
                    revision=REVISION_MAP.get(_norm(g("estado"))), gestor=GESTOR_MAP.get(_norm(g("gestor"))),
                    bloqueo=0, observacion=None, _dex=[], _fact=None)
            r = remisiones[file]
            obs = str(g("obs")).strip() if g("obs") not in (None, "") and str(g("obs")).strip() else None
            if obs and obs not in (r["observacion"] or ""):
                r["observacion"] = (r["observacion"] + " | " if r["observacion"] else "") + obs
            # ---- fechas del file
            recep = _fecha(g("recepcion"))
            borr = _fecha(g("borrador"))
            aprob = _fecha(g("aprobacion"))
            fsae = _fecha(g("fecha_sae"))
            if isinstance(g("borrador"), str) and borr:
                log["fechas_texto_convertidas"] += 1
            aprob_imp = bool(recep and aprob and aprob - recep == timedelta(days=2))
            sae_imp = bool(aprob and fsae and fsae - aprob == timedelta(hours=2, minutes=40))
            ev(recep, "H000", rb, file)
            ev(_fecha(g("clasif_sol")), "H010", rb, file)
            ev(_fecha(g("clasif_resp")), "H015", rb, file)
            ev(borr, "H020", rb, file)
            if not aprob_imp:
                ev(aprob, "H030", rb, file)
            # ---- SAE
            canal, dex_tipo = None, None
            for k in ("manifiesto", "dex"):
                t = _norm(g(k))
                if t and not re.fullmatch(r"[\d\-./ :]+", t):
                    log["texto_en_fecha"][f"{COLS[k]}: {t.lower()}"] += 1
                    if "INSPECCION" in t:
                        canal = "INSPECCION_FISICA"
                    elif "AUTORIZA" in t or "SOLICITA" in t:
                        canal = canal or "AUTORIZA_EMBARQUE"
                    if "MANUAL" in t:
                        dex_tipo = "MANUAL"
            if isinstance(g("manifiesto"), str) and _fecha(g("manifiesto")):
                log["fechas_texto_convertidas"] += 1
            dex = _codigo13(g("dex"))
            if dex:
                dex_tipo = dex_tipo or "AUTOMATICA"
            if _codigo13(g("correccion")):
                dex_tipo = "CORRECCION"
            fdex = _fecha(g("fecha_dex"))
            fman = _fecha(g("manifiesto"))
            if fman and not canal:
                canal = "AUTORIZA_EMBARQUE"
            estimado = aprob_imp or sae_imp
            if estimado:
                log["fechas_estimadas"] += 1
            causal = _norm(g("causal")) or None
            if causal and (estimado or _norm(g("cumpl_sae")) != "OUT TIME" or causal not in CAUSALES):
                causal = None
            pm, bm, pb, bu = _num(g("peso_manif")), _num(g("bultos_manif")), _num(g("bruto")), _num(g("bultos"))
            if pm is not None and pb is not None and (abs(pm - pb) > 0.5 or (bm is not None and bu is not None and bm != bu)):
                log["no_conciliadas"] += 1
            con.execute("INSERT OR REPLACE INTO sae VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
                sae_id, file, _num(g("cantidades")), _num(g("valor")), _num(g("neto")), pb,
                int(bu) if bu is not None else None, pm, int(bm) if bm is not None else None,
                int(_num(g("model")) or 0) or None, int(_num(g("partidas")) or 0) or None,
                canal, dex, dex_tipo, _codigo13(g("sae_reemp")), causal))
            con.execute("DELETE FROM evento WHERE sae = ? AND fuente = 'IMPORTACION'", (sae_id,))
            ev(fsae, "H040", rb, file, sae_id, estimado=1 if estimado else 0,
               nota="Fecha calculada por fórmula en el consolidado" if estimado else None)
            ev(fman, "H060", rb, file, sae_id)
            resp = _norm(g("responsable"))
            ev(fdex, "H070", rb, file, sae_id, resp=guardar_persona(con, resp) if resp else None)
            if fdex:
                r["_dex"].append(fdex)
            fac = _fecha(g("facturacion"))
            if fac:
                r["_fact"] = fac
            log["saes"] += 1
            log["files"].add(file)
            if rb:
                log["reembarques"].add(rb)

        for rb, d in reembarques.items():
            con.execute("INSERT OR REPLACE INTO reembarque (id, proyecto, modo, agente, naviera, terminal, pol, pod) "
                        "VALUES (?,?,?,?,?,?,?,?)", (d["id"], d["proyecto"], d["modo"], d["agente"], d["naviera"],
                                                     d["terminal"], d["pol"], d["pod"]))
        for f, r in remisiones.items():
            if r["gestor"] == "COMPLETO" and r["_dex"]:
                ev(max(r["_dex"]), "H080", r["reembarque"], f, nota="Gestor documental: DOC COMPLETOS")
            if r["_fact"]:
                ev(r["_fact"], "H090", r["reembarque"], f)
            con.execute("DELETE FROM evento WHERE file = ? AND sae IS NULL AND fuente = 'IMPORTACION'", (f,))
            con.execute("INSERT OR REPLACE INTO remision VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        [r[k] for k in ("file", "remision", "reembarque", "pais", "destinatario", "packing_list",
                                        "digitador", "apoyo", "revision", "gestor", "bloqueo")] + [None, r["observacion"]])
        con.executemany("INSERT INTO evento (fecha, nivel, reembarque, file, sae, hito, responsable, fuente, causal, nota, estimado) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?,?)", eventos)
        con.commit()
    log["eventos"] = len(eventos)
    log["files"] = len(log["files"])
    log["reembarques"] = len(log["reembarques"])
    log["texto_en_fecha"] = [{"valor": k, "n": n} for k, n in log["texto_en_fecha"].most_common(12)]
    log["transportista_desconocido"] = [{"valor": k, "n": n} for k, n in log["transportista_desconocido"].most_common()]
    log["alias_unificados"] = dict(log["alias_unificados"])
    return log


# ------------------------------------------------------------------ exportación
def exportar_excel(con):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    rems, saes = calcular(con)
    wb = Workbook()
    hojas = [
        ("Matriz", ["file", "remision", "reembarque", "pais", "destinatario", "digitador_nombre", "terminal", "naviera",
                    "contenedor", "saes", "peso_bruto", "fob_usd", "recepcion", "fase", "dueno", "accion", "ultimo",
                    "horas", "ans", "semaforo", "no_conciliadas", "revision", "gestor", "observacion"], rems),
        ("SAE", ["id", "file", "cantidad", "fob_usd", "peso_neto", "peso_bruto", "bultos", "peso_manif", "bultos_manif",
                 "dif_peso", "conciliado", "model_codes", "partidas", "canal", "dex", "dex_tipo", "h_elab", "ans_elab",
                 "h_total", "ans_sae", "causal"], saes),
        ("Eventos", ["id", "fecha", "nivel", "reembarque", "file", "sae", "hito", "responsable", "fuente", "causal",
                     "nota", "estimado"], filas(con, "SELECT * FROM evento ORDER BY fecha")),
    ]
    color = {"ROJO": "F4B6B6", "AMBAR": "FFE7A3", "VERDE": "C6EBC6", "CERRADO": "E4E4E1"}
    for i, (nombre, cols, data) in enumerate(hojas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = nombre
        ws.append([c.upper() for c in cols])
        for d in data:
            ws.append([("S" if d.get(c) else "N") if c == "conciliado" and d.get(c) is not None else d.get(c) for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
        if nombre == "Matriz":
            si = cols.index("semaforo") + 1
            for row in ws.iter_rows(min_row=2):
                v = row[si - 1].value
                if v in color:
                    row[si - 1].fill = PatternFill("solid", fgColor=color[v])
        if len(data):
            from openpyxl.utils import get_column_letter
            t = Table(displayName=f"t{nombre}", ref=f"A1:{get_column_letter(len(cols))}{len(data) + 1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
            ws.add_table(t)
        ws.freeze_panes = "B2"
        for j, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = max(10, min(40, len(c) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
